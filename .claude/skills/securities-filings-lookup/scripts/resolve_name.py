#!/usr/bin/env python3
"""
Resolve a company NAME to candidate tickers across all covered venues,
using each regulator's own lookup source. Prints one line per
candidate: venue, ticker/code, and the official name matched.

Usage:
    python resolve_name.py "Tencent"
    python resolve_name.py "Berkshire Hathaway"
    python resolve_name.py "TSMC" --venues us,tw

Sources (all free, official):
- US:      SEC's company_tickers.json (shares fetch_us_filings.py's
           one-day cache in the temp dir)
- HK:      HKEXnews prefix.do name lookup
- China:   CNINFO topSearch (matches Chinese names/pinyin; English
           names usually don't match -- expect misses for English input)
- Taiwan:  TWSE OpenAPI t187ap03_L company directory (Chinese names +
           English abbreviations; TWSE main board only)
- London:  FCA NSM search, distinct company values from recent hits

A name can legitimately resolve in several venues at once (TSMC ->
TSM ADR on EDGAR + 2330 on TWSE). Present the candidates and apply
SKILL.md's confirmation rule; don't silently pick one.
"""
from __future__ import annotations

import argparse
import io
import json
import ssl
import sys
import urllib.error
import urllib.parse
import urllib.request

import disk_cache
import sec_identity
from net_errors import HostRefused, run

if sys.stdout.encoding and sys.stdout.encoding.lower() not in ("utf-8", "utf8"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")

USER_AGENT = "Mozilla/5.0 (securities-filings-lookup-skill)"


def _context() -> ssl.SSLContext:
    try:
        import certifi
        return ssl.create_default_context(cafile=certifi.where())
    except ImportError:
        return ssl.create_default_context()


CTX = _context()


def _get(url: str, timeout: int = 30, user_agent: str | None = None) -> bytes:
    req = urllib.request.Request(
        url, headers={"User-Agent": user_agent or USER_AGENT})
    with urllib.request.urlopen(req, timeout=timeout, context=CTX) as resp:
        return resp.read()


def _post(url: str, data: bytes, content_type: str, timeout: int = 30) -> bytes:
    req = urllib.request.Request(url, data=data, headers={
        "User-Agent": USER_AGENT, "Content-Type": content_type})
    with urllib.request.urlopen(req, timeout=timeout, context=CTX) as resp:
        return resp.read()


def _cached_json(cache_name: str, url: str, ttl: int = 86400,
                 user_agent=None, inspect=None):
    """Cached JSON fetch. `user_agent` may be a callable, resolved only on
    a cache miss -- SEC needs a declared contact to make a request, not to
    read yesterday's answer off disk. `inspect` gets the raw bytes before
    they are parsed, for venues that refuse with an HTTP 200 page.

    Nothing is cached until it parses, so a refusal never lands on disk to
    be re-read for a day.
    """
    cache = disk_cache.cache_path(cache_name)
    try:
        if disk_cache.is_fresh(cache, ttl):
            with open(cache, encoding="utf-8") as f:
                return json.load(f)
    except (OSError, json.JSONDecodeError):
        pass
    raw = _get(url, user_agent=user_agent() if callable(user_agent) else user_agent)
    if inspect is not None:
        inspect(raw)
    try:
        data = json.loads(raw.decode())
    except (UnicodeDecodeError, json.JSONDecodeError) as exc:
        host = urllib.parse.urlsplit(url).hostname or url
        raise HostRefused(
            f"{host} answered with {len(raw)} bytes that are not JSON: "
            f"{raw[:80]!r}. That is the host declining this client, not an "
            "empty result.") from exc
    disk_cache.store(cache, raw)
    return data


SEC_USER_AGENT_OVERRIDE: str | None = None


def search_us(q: str) -> list[tuple[str, str, str]]:
    ql = q.lower()
    out = []
    try:
        # SEC's edge blocks browser-spoofed clients -- the module-level
        # Mozilla UA that suits the other venues is exactly what gets
        # 403'd here, so this request declares itself properly.
        data = _cached_json(
            "sec_company_tickers.json",
            "https://www.sec.gov/files/company_tickers.json",
            user_agent=lambda: sec_identity.resolve(SEC_USER_AGENT_OVERRIDE))
    except Exception as e:
        return [("us", "ERROR", str(e))]
    for entry in data.values():
        if ql in entry["title"].lower():
            out.append(("us", entry["ticker"], entry["title"]))
    return out[:8]


def search_hk(q: str) -> list[tuple[str, str, str]]:
    url = ("https://www1.hkexnews.hk/search/prefix.do?"
           + urllib.parse.urlencode({"callback": "cb", "lang": "EN",
                                     "type": "A", "name": q, "market": "SEHK"}))
    try:
        text = _get(url).decode("utf-8", errors="replace")
        inner = text[text.index("(") + 1: text.rindex(")")]
        stocks = json.loads(inner).get("stockInfo") or []
    except Exception as e:
        return [("hk", "ERROR", str(e))]
    return [("hk", s["code"] + ".HK", s["name"]) for s in stocks[:8]]


def search_cn(q: str) -> list[tuple[str, str, str]]:
    try:
        raw = _post("http://www.cninfo.com.cn/new/information/topSearch/detailOfQuery",
                    urllib.parse.urlencode({"keyWord": q, "maxSecNum": "10",
                                            "maxListNum": "5"}).encode(),
                    "application/x-www-form-urlencoded")
        entries = json.loads(raw.decode()).get("keyBoardList") or []
    except Exception as e:
        return [("cn", "ERROR", str(e))]
    return [("cn", e["code"], e.get("zwjc", "?")) for e in entries[:8]
            if e.get("category") == "A股"]


TWSE_DIRECTORY_URL = "https://openapi.twse.com.tw/v1/opendata/t187ap03_L"


def _twse_not_blocked(raw: bytes) -> None:
    """openapi.twse.com.tw serves the same HTTP-200 refusal as the document
    server, so the company directory has to recognise it too."""
    from fetch_tw_filings import is_blocked, refusal

    if is_blocked(raw.decode("utf-8", errors="replace")):
        raise refusal("openapi.twse.com.tw")


def search_tw(q: str) -> list[tuple[str, str, str]]:
    ql = q.lower()
    try:
        data = _cached_json("twse_companies.json", TWSE_DIRECTORY_URL,
                            inspect=_twse_not_blocked)
    except Exception as e:  # noqa: BLE001 - one venue failing is not fatal
        return [("tw", "ERROR", str(e))]
    out = []
    for c in data:
        names = " ".join([c.get("公司名稱", ""), c.get("公司簡稱", ""),
                          c.get("英文簡稱", "")])
        if ql in names.lower():
            out.append(("tw", c["公司代號"] + ".TW",
                        f"{c.get('公司簡稱','?')} / {c.get('英文簡稱','?')}"))
    return out[:8]


# JPX moved this list from .xls to .xlsx, which also changed which reader
# can open it: xlrd 2.x dropped xlsx support entirely, so openpyxl it is.
# The old .xls URL now 404s -- and that went unnoticed because the missing
# xlrd was reported first, so the venue said "SKIPPED" rather than "the
# directory has moved".
JPX_DIRECTORY_URL = ("https://www.jpx.co.jp/english/markets/statistics-equities/"
                     "misc/tvdivq0000001vg2-att/data_e.xlsx")
# Columns, confirmed against the live file (4,441 rows, 2026-09):
# 0 Effective Date | 1 Local Code | 2 Name (English) | 3 Section/Products
JPX_CODE_COLUMN = 1
JPX_NAME_COLUMN = 2


def search_jp(q: str) -> list[tuple[str, str, str]]:
    """JPX's English listed-company directory (an .xlsx; needs openpyxl)."""
    try:
        import openpyxl
    except ImportError:
        return [("jp", "SKIPPED",
                 "pip install openpyxl to enable Japan name lookup")]
    cache = disk_cache.cache_path("jpx_companies_e.xlsx")
    try:
        if disk_cache.is_fresh(cache):
            source = cache
        else:
            data = _get(JPX_DIRECTORY_URL, timeout=90)
            disk_cache.store(cache, data)
            # Read the bytes in hand, not the file just written: store()
            # swallows write failures by design, and reading back a cache
            # that never landed would fail the lookup that had already
            # succeeded.
            source = io.BytesIO(data)
        book = openpyxl.load_workbook(source, read_only=True, data_only=True)
    except Exception as e:  # noqa: BLE001 - one venue failing is not fatal
        return [("jp", "ERROR", str(e))]
    ql = q.lower()
    out = []
    try:
        sheet = book[book.sheetnames[0]]
        for row in sheet.iter_rows(min_row=2, values_only=True):
            name = str(row[JPX_NAME_COLUMN] or "")
            if ql and ql in name.lower():
                code = str(row[JPX_CODE_COLUMN] or "").split(".")[0]
                out.append(("jp", code + ".T", name))
                if len(out) >= 8:
                    break
    except Exception as e:  # noqa: BLE001 - a layout change, not a crash
        return [("jp", "ERROR", f"JPX directory layout changed: {e}")]
    finally:
        book.close()
    return out


def search_uk(q: str) -> list[tuple[str, str, str]]:
    body = {"from": 0, "size": 40, "sort": "publication_date", "sortorder": "desc",
            "keyword": q, "criteriaObj": {"criteria": [], "dateCriteria": []}}
    # The endpoint and the meaning of a 400 from it are defined once, in
    # fetch_uk_filings, so the next FCA change is a single edit.
    from fetch_uk_filings import RETIRED_INDEX_NOTE, SEARCH_URL
    try:
        raw = _post(SEARCH_URL, json.dumps(body).encode(), "application/json")
        hits = json.loads(raw.decode()).get("hits", {}).get("hits", [])
    except urllib.error.HTTPError as e:
        if e.code == 400:
            return [("uk", "ERROR", RETIRED_INDEX_NOTE)]
        return [("uk", "ERROR", str(e))]
    except Exception as e:
        return [("uk", "ERROR", str(e))]
    seen, out = set(), []
    ql = q.lower()
    for h in hits:
        comp = (h["_source"].get("company") or "").strip("; ")
        sym = h["_source"].get("symbol") or "?"
        if comp and comp not in seen and ql.split()[0] in comp.lower():
            seen.add(comp)
            out.append(("uk", sym, comp))
    return out[:8]


def search_sg(q: str) -> list[tuple[str, str, str]]:
    """SGX's own securities directory: code, name and security type.

    Reporting issuers are listed first -- a search for "DBS" should lead
    with the bank, not with one of its leveraged certificates.
    """
    try:
        from fetch_sg_filings import find, is_issuer, load_directory
        matches = find(q, load_directory())
    except Exception as e:  # noqa: BLE001 - one venue failing is not fatal
        return [("sg", "ERROR", str(e))]
    out = []
    for row in matches:
        label = row["name"] + ("" if is_issuer(row) else f" ({row['type']})")
        out.append(("sg", row["code"] + ".SI", label))
    return out[:8]


VENUES = {"us": search_us, "hk": search_hk, "cn": search_cn,
          "tw": search_tw, "uk": search_uk, "jp": search_jp,
          "sg": search_sg}


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("name", help="Company name (or fragment) to resolve")
    parser.add_argument("--venues", default="us,hk,cn,tw,uk,jp,sg",
                        help="Comma-separated subset of us,hk,cn,tw,uk,jp,sg "
                             "(no directory exists for Germany -- web search)")
    parser.add_argument("--user-agent",
                        help="Contact for SEC's fair-access policy, e.g. "
                             f"'{sec_identity.EXAMPLE_CONTACT}'")
    args = parser.parse_args()

    global SEC_USER_AGENT_OVERRIDE
    SEC_USER_AGENT_OVERRIDE = args.user_agent

    any_hit = False
    for v in args.venues.split(","):
        fn = VENUES.get(v.strip())
        if not fn:
            continue
        for venue, code, name in fn(args.name):
            print(f"{venue:<3} {code:<12} {name}")
            # An ERROR or SKIPPED row is a venue that could not answer, not
            # a candidate -- counting it suppresses the fallback advice.
            if code not in ("ERROR", "SKIPPED"):
                any_hit = True
    if not any_hit:
        print("No candidates found in any venue directory. Fall back to a web "
              "search for '<name> stock ticker' and confirm with the user.")


if __name__ == "__main__":
    run(main)
